"""Genera los Excel de ejemplo a partir de sus CSV.

- samples/invitados.xlsx           lista de prueba (18 invitados, con errores a propósito)
- samples/plantilla-invitados.xlsx plantilla para que la organizadora complete y suba

Los teléfonos se escriben como TEXTO a propósito: si Excel los interpreta como
número, pierde el cero inicial ("0981 123 456") y la importación rechaza la fila.
"""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="6B5B4A")
WIDTHS = [28, 20, 8, 24, 34]

INSTRUCCIONES = [
    ("Columna", "Qué poner", "¿Obligatoria?"),
    (
        "Nombre y Apellido",
        "El titular de la invitación. Una fila por invitación, no una por persona: "
        "los acompañantes se cargan al confirmar.",
        "Sí",
    ),
    (
        "Celular",
        "El número de WhatsApp. Puede ir con 0 inicial, con +595 o con guiones: "
        "el sistema lo normaliza solo.",
        "Sí",
    ),
    (
        "Pases",
        "Cuántas personas entran con esta invitación: 1, 2, 4… Si se deja vacío se usan 2.",
        "No",
    ),
    (
        "Familia",
        "Grupo, familia o mesa. Sirve para agrupar y para enviar un solo mensaje por familia.",
        "No",
    ),
    (
        "Observaciones",
        "Restricciones alimentarias, si vienen niños, cualquier dato útil para el salón.",
        "No",
    ),
    ("", "", ""),
    (
        "Cómo subirla",
        "Panel → Invitados → Importar → arrastrar este archivo. Se muestra una vista previa "
        "con los errores fila por fila antes de guardar nada.",
        "",
    ),
    ("Consejo", "Borrá las filas de ejemplo antes de subir tu lista real.", ""),
]


def build(csv_name: str, xlsx_name: str, sheet_title: str, with_instructions: bool) -> None:
    csv_path = BASE / csv_name
    with csv_path.open(encoding="utf-8", newline="") as fh:
        rows = list(csv.reader(fh))

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if r == 1:
                cell.font = HEADER_FONT
                cell.fill = HEADER_FILL
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.number_format = "@"  # texto: conserva ceros y espacios
                cell.alignment = Alignment(vertical="center")

    for i, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(rows[0]))}{len(rows)}"

    if with_instructions:
        help_sheet = wb.create_sheet("Instrucciones")
        for r, row in enumerate(INSTRUCCIONES, start=1):
            for c, value in enumerate(row, start=1):
                cell = help_sheet.cell(row=r, column=c, value=value)
                if r == 1:
                    cell.font = HEADER_FONT
                    cell.fill = HEADER_FILL
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        help_sheet.column_dimensions["A"].width = 22
        help_sheet.column_dimensions["B"].width = 78
        help_sheet.column_dimensions["C"].width = 16

    out = BASE / xlsx_name
    wb.save(out)
    print(f"OK -> {out}  ({len(rows) - 1} filas de datos, hojas: {', '.join(wb.sheetnames)})")


build("invitados.csv", "invitados.xlsx", "Invitados", with_instructions=False)
build("plantilla-invitados.csv", "plantilla-invitados.xlsx", "Invitados", with_instructions=True)
