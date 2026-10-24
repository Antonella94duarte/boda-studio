"""Genera el manual de la organizadora en DOCX a partir del Markdown.

Uso:
    python scripts/build-manual-docx.py

Entrada:  docs/08-manual-organizadora.md  (con sus imágenes en docs/img/manual/)
Salida:   docs/Boda-Studio-manual-de-la-organizadora.docx

Es un conversor del subconjunto de Markdown que usa el manual: títulos,
párrafos, listas, tablas, imágenes con pie, citas, bloques de código y líneas
horizontales. Así el manual tiene una sola fuente de verdad (el .md del repo) y
el Word se regenera cuando cambia.
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

BRAND = RGBColor(0x3B, 0x33, 0x2A)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
INK = RGBColor(0x4A, 0x40, 0x38)
MUTED = RGBColor(0x8A, 0x7A, 0x63)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

BODY_FONT = "Calibri"
MONO_FONT = "Consolas"
TEXT_WIDTH_CM = 16.4

REPO = Path(__file__).resolve().parent.parent
SOURCE = REPO / "docs" / "08-manual-organizadora.md"
OUT = REPO / "docs" / "Boda-Studio-manual-de-la-organizadora.docx"


# --------------------------------------------------------------------------- #
# Utilidades de formato
# --------------------------------------------------------------------------- #
def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def cell_margins(cell, top=70, bottom=70, left=100, right=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag, value in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


def table_borders(table, color="D8D0C4", size=4):
    tbl_pr = table._tbl.tblPr
    for existing in tbl_pr.findall(qn("w:tblBorders")):
        tbl_pr.remove(existing)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def fix_widths(table, widths_cm):
    tbl_pr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        for existing in tbl_pr.findall(qn(tag)):
            tbl_pr.remove(existing)

    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(int(round(sum(widths_cm) * 567))))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)

    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)

    for index, width in enumerate(widths_cm):
        table.columns[index].width = Cm(width)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_field(paragraph, instruction, placeholder="1"):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, sep, text, end):
        run._r.append(element)


def spacer(doc, pts=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1
    run = p.add_run("")
    run.font.size = Pt(pts)


def horizontal_rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(10)
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "E7E1D8")
    borders.append(bottom)
    p_pr.append(borders)


# --------------------------------------------------------------------------- #
# Texto en línea: **negrita**, `código`
# --------------------------------------------------------------------------- #
INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def add_inline(paragraph, text, size=10.5, color=INK, bold_all=False):
    for chunk in INLINE_RE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            run = paragraph.add_run(chunk[2:-2])
            run.font.bold = True
            run.font.name = BODY_FONT
            run.font.size = Pt(size)
            run.font.color.rgb = BRAND
        elif chunk.startswith("`") and chunk.endswith("`") and len(chunk) > 2:
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = MONO_FONT
            run.font.size = Pt(max(8.0, size - 1.5))
            run.font.color.rgb = color
        else:
            run = paragraph.add_run(chunk)
            run.font.bold = bold_all
            run.font.name = BODY_FONT
            run.font.size = Pt(size)
            run.font.color.rgb = color


# --------------------------------------------------------------------------- #
# Bloques
# --------------------------------------------------------------------------- #
def heading(doc, text, level):
    if level == 1:
        kicker = doc.add_paragraph()
        kicker.paragraph_format.space_after = Pt(2)
        run = kicker.add_run("MANUAL DE LA ORGANIZADORA")
        run.font.name = BODY_FONT
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = GOLD

        title = doc.add_paragraph()
        title.paragraph_format.space_after = Pt(10)
        run = title.add_run(text.replace("08 · ", ""))
        run.font.name = BODY_FONT
        run.font.size = Pt(26)
        run.font.bold = True
        run.font.color.rgb = BRAND
        return

    h = doc.add_heading(text, level=1)
    h.paragraph_format.space_before = Pt(16 if level == 2 else 12)
    h.paragraph_format.space_after = Pt(5)
    h.paragraph_format.keep_with_next = True
    for run in h.runs:
        run.font.name = BODY_FONT
        run.font.size = Pt(13.5 if level == 2 else 11.5)
        run.font.bold = True
        run.font.color.rgb = BRAND


def paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(p, text)


def bullet(doc, text, number=None):
    p = doc.add_paragraph(style="List Bullet" if number is None else "List Number")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    add_inline(p, text)
    return p


def quote(doc, text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.autofit = False
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(tbl, color="E3DACA", size=4)
    fix_widths(tbl, [TEXT_WIDTH_CM])
    cell = tbl.rows[0].cells[0]
    cell.width = Cm(TEXT_WIDTH_CM)
    shade(cell, "FDF1D8")
    cell_margins(cell, top=110, bottom=110, left=140, right=140)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    add_inline(p, text, size=9.5, color=RGBColor(0x6B, 0x47, 0x00))
    spacer(doc, 8)


def code_block(doc, lines):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.autofit = False
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(tbl, color="E3DACA", size=4)
    fix_widths(tbl, [TEXT_WIDTH_CM])
    cell = tbl.rows[0].cells[0]
    cell.width = Cm(TEXT_WIDTH_CM)
    shade(cell, "F7F4EF")
    cell_margins(cell, top=110, bottom=110, left=140, right=140)
    first = True
    for line in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(line)
        run.font.name = MONO_FONT
        run.font.size = Pt(9)
        run.font.color.rgb = INK
    spacer(doc, 8)


def image(doc, path, caption):
    target = (SOURCE.parent / path).resolve()
    if not target.exists():
        print(f"  aviso: falta la imagen {path}")
        return

    with Image.open(target) as probe:
        width_px, height_px = probe.size

    max_width = 13.0
    max_height = 15.5
    width_cm = min(max_width, width_px / 96 * 2.54)
    height_cm = width_cm * height_px / width_px
    # Las capturas de la invitación son de teléfono (muy verticales): sin este
    # tope ocuparían una página entera.
    if height_cm > max_height:
        height_cm = max_height
        width_cm = height_cm * width_px / height_px

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(target), width=Cm(width_cm), height=Cm(height_cm))

    if caption:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(10)
        run = c.add_run(caption)
        run.font.name = BODY_FONT
        run.font.size = Pt(8.5)
        run.font.italic = True
        run.font.color.rgb = MUTED


def table(doc, rows):
    header = rows[0]
    body = rows[1:]

    # Ancho proporcional al texto de cada columna, con un mínimo legible.
    weights = []
    for index in range(len(header)):
        longest = max(
            [len(header[index].strip())] + [len(row[index].strip()) for row in body if index < len(row)]
        )
        weights.append(max(12, min(longest, 60)))
    total = sum(weights)
    widths = [max(2.4, TEXT_WIDTH_CM * weight / total) for weight in weights]
    scale = TEXT_WIDTH_CM / sum(widths)
    widths = [round(width * scale, 2) for width in widths]

    tbl = doc.add_table(rows=1, cols=len(header))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    table_borders(tbl)
    fix_widths(tbl, widths)

    head = tbl.rows[0]
    repeat_header(head)
    for index, text in enumerate(header):
        cell = head.cells[index]
        cell.width = Cm(widths[index])
        shade(cell, "3B332A")
        cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_inline(p, text, size=9, color=WHITE, bold_all=True)

    for row_index, values in enumerate(body):
        cells = tbl.add_row().cells
        for col_index in range(len(header)):
            value = values[col_index] if col_index < len(values) else ""
            cell = cells[col_index]
            cell.width = Cm(widths[col_index])
            cell_margins(cell)
            if row_index % 2 == 1:
                shade(cell, "F7F4EF")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            add_inline(p, value, size=9)

    spacer(doc, 8)


# --------------------------------------------------------------------------- #
# Conversión
# --------------------------------------------------------------------------- #
def split_row(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)

    base = doc.styles["Normal"]
    base.font.name = BODY_FONT
    base.font.size = Pt(10.5)
    base.font.color.rgb = INK
    base.paragraph_format.space_after = Pt(6)

    index = 0
    paragraph_buffer = []

    def flush_paragraph():
        nonlocal paragraph_buffer
        if paragraph_buffer:
            paragraph(doc, " ".join(paragraph_buffer))
            paragraph_buffer = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_paragraph()
            index += 1
            code = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            index += 1
            code_block(doc, code)
            continue

        if stripped == "---":
            flush_paragraph()
            horizontal_rule(doc)
            index += 1
            continue

        if not stripped:
            flush_paragraph()
            index += 1
            continue

        if stripped.startswith("#"):
            flush_paragraph()
            level = len(stripped) - len(stripped.lstrip("#"))
            heading(doc, stripped[level:].strip(), level)
            index += 1
            continue

        image_match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", stripped)
        if image_match:
            flush_paragraph()
            image(doc, image_match.group(2), image_match.group(1))
            index += 1
            continue

        if stripped.startswith("|"):
            flush_paragraph()
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                row = split_row(lines[index])
                # Fila separadora de Markdown: |---|---|
                if not all(re.fullmatch(r":?-{2,}:?", cell or "-") for cell in row):
                    rows.append(row)
                index += 1
            if rows:
                table(doc, rows)
            continue

        if stripped.startswith("> "):
            flush_paragraph()
            quote_lines = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip().lstrip("> ").strip())
                index += 1
            quote(doc, " ".join(quote_lines))
            continue

        ordered = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if ordered:
            flush_paragraph()
            bullet(doc, ordered.group(2), number=int(ordered.group(1)))
            index += 1
            continue

        if stripped.startswith("- "):
            flush_paragraph()
            bullet(doc, stripped[2:].strip())
            index += 1
            continue

        paragraph_buffer.append(stripped)
        index += 1

    flush_paragraph()

    # Pie de página con numeración.
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("Boda Studio · Manual de la organizadora · página ")
    add_field(fp, "PAGE")
    run = fp.add_run(" de ")
    add_field(fp, "NUMPAGES", placeholder="20")
    for run in fp.runs:
        run.font.name = BODY_FONT
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"OK -> {OUT}")


if __name__ == "__main__":
    build()
